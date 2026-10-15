import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

COLOR_PRIMARY = RGBColor(30, 58, 138)     # Deep Navy #1E3A8A
COLOR_SECONDARY = RGBColor(15, 118, 110)  # Deep Teal #0F766E
COLOR_TEXT = RGBColor(30, 41, 59)         # Slate 800 #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate 500 #64748B
COLOR_ACCENT = RGBColor(180, 83, 9)       # Amber 700 #B45309

HEX_PRIMARY = "1E3A8A"
HEX_LIGHT_BG = "F1F5F9"
HEX_ROW_ALT = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_SHORTCUT_BG = "F0FDF4"
HEX_SHORTCUT_BORDER = "16A34A"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shd if any
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def init_document():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
    
    # Configure Normal style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = COLOR_TEXT
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15
    return doc

def add_header_block(doc, unit_no, unit_name, area_name):
    # Top subtitle / series
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(2)
    run_meta = p_meta.add_run(f"Mathematics XI  |  NEB (Grade 11)  |  Mat. 007  |  Area: {area_name}")
    run_meta.font.size = Pt(9.5)
    run_meta.font.bold = True
    run_meta.font.color.rgb = COLOR_MUTED

    # Unit Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(f"UNIT {unit_no}: {unit_name.upper()}")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    # Description line
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(0)
    p_desc.paragraph_format.space_after = Pt(14)
    run_desc = p_desc.add_run("All-possible NEB questions  •  short precise answers  •  conceptual insights  •  exam shortcuts  •  worked examples")
    run_desc.font.size = Pt(10)
    run_desc.font.italic = True
    run_desc.font.color.rgb = COLOR_SECONDARY

def add_topic_heading(doc, topic_code, topic_title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run_symbol = p.add_run("▣ ")
    run_symbol.font.bold = True
    run_symbol.font.size = Pt(13)
    run_symbol.font.color.rgb = COLOR_PRIMARY

    run = p.add_run(f"{topic_code} {topic_title}")
    run.font.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = COLOR_PRIMARY

def add_question(doc, q_num, question_text, answer_text, shortcut_text=None):
    # Question line
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(7)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.keep_with_next = True

    run_q_label = p_q.add_run(f"Q{q_num}. ")
    run_q_label.font.bold = True
    run_q_label.font.color.rgb = COLOR_PRIMARY

    run_q_text = p_q.add_run(question_text)
    run_q_text.font.bold = True
    run_q_text.font.color.rgb = COLOR_TEXT

    # Answer line
    p_a = doc.add_paragraph()
    p_a.paragraph_format.space_before = Pt(0)
    p_a.paragraph_format.space_after = Pt(2)
    
    run_a_label = p_a.add_run("Answer:  ")
    run_a_label.font.bold = True
    run_a_label.font.color.rgb = RGBColor(71, 85, 105) # Slate 600

    run_a_text = p_a.add_run(answer_text)
    run_a_text.font.color.rgb = COLOR_TEXT

    # Shortcut line if provided
    if shortcut_text:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(0)
        p_s.paragraph_format.space_after = Pt(6)
        p_s.paragraph_format.left_indent = Inches(0.2)
        
        run_s_icon = p_s.add_run("⚡ Shortcut / Exam Insight:  ")
        run_s_icon.font.bold = True
        run_s_icon.font.size = Pt(9.5)
        run_s_icon.font.color.rgb = COLOR_ACCENT

        run_s_text = p_s.add_run(shortcut_text)
        run_s_text.font.size = Pt(9.5)
        run_s_text.font.italic = True
        run_s_text.font.color.rgb = RGBColor(67, 56, 202) # Indigo 700

def add_compare_section_header(doc, topic_code, topic_title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True

    run = p.add_run(f"▥ {topic_code} {topic_title}  —  Compare / Contrast Tables")
    run.font.bold = True
    run.font.size = Pt(11.5)
    run.font.color.rgb = COLOR_SECONDARY

def add_compare_table(doc, table_title, headers, rows):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(3)
    p_title.paragraph_format.keep_with_next = True

    run_title = p_title.add_run(table_title)
    run_title.font.bold = True
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = COLOR_PRIMARY

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=HEX_BORDER, sz="4")

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], HEX_PRIMARY)
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=120, right=120)

    # Data Rows
    for r_idx, r_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = HEX_ROW_ALT if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, cell_value in enumerate(r_data):
            row_cells[c_idx].text = ""
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(cell_value))
            run.font.size = Pt(9.5)
            run.font.color.rgb = COLOR_TEXT
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=120, right=120)

    # Space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def add_viva_header(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True

    run = p.add_run("PROFESSOR'S THEORY & VIVA BANK   — the kinds of theory a professor / examiner actually asks")
    run.font.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = COLOR_PRIMARY

def add_viva_category(doc, category_name, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True

    run = p.add_run(f"▹ {category_name}  ({subtitle})")
    run.font.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_SECONDARY

def add_viva_question(doc, category, v_num, question_text, answer_text):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(5)
    p_q.paragraph_format.space_after = Pt(2)
    p_q.paragraph_format.keep_with_next = True

    run_tag = p_q.add_run(f"[{category}] ")
    run_tag.font.bold = True
    run_tag.font.color.rgb = COLOR_ACCENT

    run_num = p_q.add_run(f"T{v_num}. ")
    run_num.font.bold = True
    run_num.font.color.rgb = COLOR_PRIMARY

    run_text = p_q.add_run(question_text)
    run_text.font.bold = True
    run_text.font.color.rgb = COLOR_TEXT

    p_a = doc.add_paragraph()
    p_a.paragraph_format.space_before = Pt(0)
    p_a.paragraph_format.space_after = Pt(4)

    run_lbl = p_a.add_run("Answer / Proof sketch:  ")
    run_lbl.font.bold = True
    run_lbl.font.color.rgb = RGBColor(71, 85, 105)

    run_ans = p_a.add_run(answer_text)
    run_ans.font.color.rgb = COLOR_TEXT

def add_quick_memos_section(doc, memos):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True

    run = p.add_run("★ Quick-math memos (this unit)")
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = COLOR_ACCENT

    for memo in memos:
        p_m = doc.add_paragraph(style='List Bullet')
        p_m.paragraph_format.space_before = Pt(1)
        p_m.paragraph_format.space_after = Pt(2)
        run_m = p_m.add_run(memo)
        run_m.font.size = Pt(10)
        run_m.font.color.rgb = COLOR_TEXT
